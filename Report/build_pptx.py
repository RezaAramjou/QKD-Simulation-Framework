#!/usr/bin/env python3
"""build_pptx.py — Parse PRESENTATION_DECK.md and produce Presentation.pptx"""

import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BG_DARK    = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT     = RGBColor(0x00, 0xB4, 0xD8)
ACCENT2    = RGBColor(0x48, 0xCA, 0xE4)
GOLD       = RGBColor(0xFF, 0xD1, 0x66)
LIGHT_TEXT  = RGBColor(0xF0, 0xF0, 0xF0)
SUBTLE      = RGBColor(0xAA, 0xAA, 0xAA)
SLIDE_NUM   = RGBColor(0x90, 0xAD, 0xC2)

def sanitise(text):
    text = text.replace('**', '')
    text = text.replace('__', '')
    text = text.replace('*', '')
    return re.sub(r'\s+', ' ', text).strip()

def parse_md(path):
    raw = Path(path).read_text(encoding='utf-8')
    lines = raw.split('\n')
    meta = {}
    slides = []
    cur = None
    section = None
    buf = []
    in_header = True
    in_time = False
    title_line = ''

    def flush():
        nonlocal buf
        if section and buf:
            txt = '\n'.join(buf).strip()
            if section == 'script':
                txt = txt.strip().strip('"').strip("'").strip()
            if cur is not None:
                cur[section] = txt
        buf.clear()

    for line in lines:
        if in_header:
            if line.startswith('# ') and not line.startswith('## '):
                title_line = line[2:].strip()
                continue
            if line.startswith('**Presenter:**'):
                meta['presenter'] = sanitise(line)
                continue
            if line.startswith('**Duration:**'):
                meta['duration'] = sanitise(line)
                continue
            if line.startswith('**Date:**'):
                meta['date'] = sanitise(line)
                continue
            if line.strip().startswith('## Time Allocation'):
                in_time = True
                continue
            if in_time:
                if line.strip().startswith('|'):
                    continue
                in_time = False
            if line.strip() == '---':
                in_header = False
                continue
            continue

        if line.strip().startswith('## SLIDE') or (line.strip().startswith('**[Slide') and cur is None):
            flush()
            cur = {}
            section = None
            m = re.match(r'##\s+SLIDE\s+(\d+)', line.strip())
            if m:
                cur['num'] = int(m.group(1))
            continue

        if line.strip().startswith('**[Slide'):
            flush()
            if cur is None:
                cur = {}
            m = re.match(r'\*\*\[Slide\s+(\d+)\]\:', line.strip())
            if m:
                cur['num'] = int(m.group(1))
            title = re.sub(r'^\*\*\[Slide\s+\d+\]:\s*', '', line.strip()).rstrip('*').strip()
            cur['title'] = title
            continue

        if line.strip() == '---':
            flush()
            if cur:
                slides.append(cur)
            cur = None
            section = None
            continue

        if cur is None:
            continue

        if line.strip().startswith('- **Visuals:**'):
            flush()
            section = 'visuals'
            buf.append(line.strip()[len('- **Visuals:**'):].strip())
            continue
        if line.strip().startswith('- **Slide Text:**'):
            flush()
            section = 'text'
            buf.append(line.strip()[len('- **Slide Text:**'):].strip())
            continue
        if line.strip().startswith('- **Script:**'):
            flush()
            section = 'script'
            buf.append(line.strip()[len('- **Script:**'):].strip())
            continue

        if section:
            buf.append(line)

    flush()
    if cur:
        slides.append(cur)
    meta['title'] = title_line or 'QKD Simulation Framework'
    return meta, slides


def add_dark_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

def add_accent_bar(slide, prs, top=0, height=Inches(0.06)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, prs.slide_width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()

def add_slide_number(slide, num, prs):
    txBox = slide.shapes.add_textbox(prs.slide_width - Inches(0.8), prs.slide_height - Inches(0.5), Inches(0.6), Inches(0.4))
    p = txBox.text_frame.paragraphs[0]
    p.text = str(num)
    p.font.size = Pt(10)
    p.font.color.rgb = SLIDE_NUM
    p.alignment = PP_ALIGN.RIGHT

def add_section_tag(slide, tag, prs):
    txBox = slide.shapes.add_textbox(prs.slide_width - Inches(2.2), Inches(0.15), Inches(2.0), Inches(0.3))
    p = txBox.text_frame.paragraphs[0]
    p.text = tag.upper()
    p.font.size = Pt(8)
    p.font.color.rgb = ACCENT
    p.font.bold = True
    p.alignment = PP_ALIGN.RIGHT


def build_title_slide(prs, meta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_dark_bg(slide)
    add_accent_bar(slide, prs, top=Inches(0), height=Inches(0.08))
    add_accent_bar(slide, prs, top=prs.slide_height - Inches(0.08), height=Inches(0.08))

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8.4), Inches(1.4))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = meta.get('title', 'QKD Simulation Framework')
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.font.bold = True

    txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(8.4), Inches(1.0))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = 'A Unified Simulation Framework for Discrete-Variable\nQuantum Key Distribution'
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT2

    txBox3 = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(8.4), Inches(1.2))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    for i, label in enumerate([meta.get('presenter',''), meta.get('date',''), meta.get('duration','')]):
        p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        p.text = label
        p.font.size = Pt(14)
        p.font.color.rgb = SUBTLE
        p.space_after = Pt(6)
    return slide


def build_content_slide(prs, sd):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_dark_bg(slide)
    add_accent_bar(slide, prs, top=Inches(0), height=Inches(0.05))

    num = sd.get('num', 0)
    title = sd.get('title', '')
    text_body = sd.get('text', '')
    script = sd.get('script', '')

    add_section_tag(slide, 'Slide %d' % num, prs)
    add_slide_number(slide, num, prs)

    txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.color.rgb = GOLD
    p.font.bold = True

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.15), Inches(2.5), Inches(0.03))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()

    bullets = [b.strip().lstrip('- ') for b in text_body.split('\n') if b.strip()]
    txBox2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(8.8), Inches(3.0))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = bullet
        p.font.size = Pt(15)
        p.font.color.rgb = LIGHT_TEXT
        p.space_after = Pt(8)

    if script:
        short = script[:400] + '\u2026' if len(script) > 400 else script
        txBox3 = slide.shapes.add_textbox(Inches(0.6), Inches(4.8), Inches(8.8), Inches(2.0))
        tf3 = txBox3.text_frame
        tf3.word_wrap = True
        p3 = tf3.paragraphs[0]
        p3.text = '\U0001F3A4 ' + short
        p3.font.size = Pt(10)
        p3.font.color.rgb = SUBTLE
        p3.font.italic = True

    return slide


def main():
    md_path = Path(__file__).parent / 'PRESENTATION_DECK.md'
    meta, slides = parse_md(str(md_path))
    print('Parsed %d slides' % len(slides))
    for s in slides:
        print('  Slide %s: %s' % (s.get('num','?'), s.get('title','NO TITLE')[:60]))

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    build_title_slide(prs, meta)
    for sd in slides:
        build_content_slide(prs, sd)

    out = Path(__file__).parent / 'Presentation.pptx'
    prs.save(str(out))
    print('\nSaved %s (%d slides)' % (out, len(slides)+1))

if __name__ == '__main__':
    main()
